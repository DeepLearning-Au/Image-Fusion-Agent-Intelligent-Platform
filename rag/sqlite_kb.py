from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
import re
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional

try:
    from config import settings
except Exception:
    settings = None

PROJECT_DIR = Path(__file__).resolve().parents[1]
STORAGE_DIR = PROJECT_DIR / "storage"
DB_PATH = STORAGE_DIR / "kb.sqlite3"
STORAGE_DIR.mkdir(parents=True, exist_ok=True)

SUPPORTED_EXTS = {".txt", ".md", ".csv", ".py", ".yaml", ".yml", ".json", ".html", ".htm", ".docx", ".pdf"}


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


@contextmanager
def connect() -> Iterator[sqlite3.Connection]:
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA busy_timeout = 5000")
    conn.execute("PRAGMA journal_mode = WAL")
    try:
        yield conn
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def has_fts_table(conn: sqlite3.Connection) -> bool:
    row = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'chunks_fts'"
    ).fetchone()
    return row is not None


def ensure_bm25_fts_table(conn: sqlite3.Connection) -> None:
    """Create the FTS5 index used by SQLite's native BM25 ranker.

    Trigram tokenization makes continuous Chinese text searchable without an
    additional segmentation service. Existing FTS5 tables remain compatible
    with native BM25 even when an older SQLite build lacks trigram support.
    """
    if not has_fts_table(conn):
        try:
            conn.execute("""
            CREATE VIRTUAL TABLE chunks_fts USING fts5(
                title,
                category,
                content,
                doc_id UNINDEXED,
                chunk_id UNINDEXED,
                tokenize='trigram'
            )
            """)
        except sqlite3.OperationalError:
            # Older SQLite builds may not include the trigram tokenizer. FTS5
            # BM25 remains available; LIKE is retained as the Chinese fallback.
            conn.execute("""
            CREATE VIRTUAL TABLE chunks_fts USING fts5(
                title,
                category,
                content,
                doc_id UNINDEXED,
                chunk_id UNINDEXED
            )
            """)

    indexed_count = int(conn.execute("SELECT COUNT(*) FROM chunks_fts").fetchone()[0])
    chunk_count = int(conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0])
    if indexed_count != chunk_count:
        conn.execute("DELETE FROM chunks_fts")
        conn.execute("""
            INSERT INTO chunks_fts(title, category, content, doc_id, chunk_id)
            SELECT d.title, d.category, c.content, c.doc_id, c.id
            FROM chunks c
            JOIN documents d ON d.id = c.doc_id
            ORDER BY c.id
        """)


def table_columns(conn: sqlite3.Connection, table: str) -> set:
    rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    return {row[1] for row in rows}


def reset_schema(conn: sqlite3.Connection) -> None:
    for table in ["chunk_embeddings", "document_tables", "document_pages", "chunks_fts", "chunks", "parent_chunks", "documents", "kb_generation_answer_evaluations", "kb_generation_evaluations", "kb_generation_events", "kb_generations"]:
        conn.execute(f"DROP TABLE IF EXISTS {table}")


def ensure_column(conn: sqlite3.Connection, table: str, column: str, definition: str) -> None:
    if column not in table_columns(conn, table):
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def ensure_compatible_schema(conn: sqlite3.Connection) -> None:
    doc_cols = table_columns(conn, "documents")
    chunk_cols = table_columns(conn, "chunks")
    if doc_cols and not {"id", "title", "source", "category", "enabled"}.issubset(doc_cols):
        reset_schema(conn)
        return
    if chunk_cols and not {"id", "doc_id", "chunk_index", "content"}.issubset(chunk_cols):
        reset_schema(conn)


def init_db() -> None:
    with connect() as conn:
        ensure_compatible_schema(conn)
        conn.execute("""
        CREATE TABLE IF NOT EXISTS kb_generations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            status TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            published_at TEXT DEFAULT '',
            error_message TEXT DEFAULT '',
            metadata_json TEXT DEFAULT '{}'
        )
        """)
        conn.execute(
            "CREATE UNIQUE INDEX IF NOT EXISTS idx_kb_one_active_generation "
            "ON kb_generations(status) WHERE status = 'active'"
        )
        conn.execute("""
        CREATE TABLE IF NOT EXISTS kb_generation_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            generation_id INTEGER NOT NULL,
            event_type TEXT NOT NULL,
            detail TEXT DEFAULT '',
            created_at TEXT NOT NULL,
            FOREIGN KEY(generation_id) REFERENCES kb_generations(id) ON DELETE CASCADE
        )
        """)
        conn.execute("""
        CREATE TABLE IF NOT EXISTS kb_generation_evaluations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            generation_id INTEGER NOT NULL UNIQUE,
            eval_set_path TEXT DEFAULT '',
            total INTEGER NOT NULL DEFAULT 0,
            recall_at_k REAL NOT NULL DEFAULT 0,
            mrr REAL NOT NULL DEFAULT 0,
            ndcg_at_k REAL NOT NULL DEFAULT 0,
            top1_accuracy REAL NOT NULL DEFAULT 0,
            keyword_coverage REAL NOT NULL DEFAULT 0,
            passed INTEGER NOT NULL DEFAULT 0,
            thresholds_json TEXT NOT NULL DEFAULT '{}',
            report_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL,
            FOREIGN KEY(generation_id) REFERENCES kb_generations(id) ON DELETE CASCADE
        )
        """)
        conn.execute("""
        CREATE TABLE IF NOT EXISTS kb_generation_answer_evaluations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            generation_id INTEGER NOT NULL UNIQUE,
            eval_set_path TEXT DEFAULT '',
            total INTEGER NOT NULL DEFAULT 0,
            answer_relevance REAL NOT NULL DEFAULT 0,
            faithfulness REAL NOT NULL DEFAULT 0,
            citation_completeness REAL NOT NULL DEFAULT 0,
            refusal_correctness REAL NOT NULL DEFAULT 0,
            passed INTEGER NOT NULL DEFAULT 0,
            thresholds_json TEXT NOT NULL DEFAULT '{}',
            report_json TEXT NOT NULL DEFAULT '{}',
            created_at TEXT NOT NULL,
            FOREIGN KEY(generation_id) REFERENCES kb_generations(id) ON DELETE CASCADE
        )
        """)
        conn.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            generation_id INTEGER,
            title TEXT NOT NULL,
            source TEXT DEFAULT '',
            category TEXT DEFAULT '',
            enabled INTEGER DEFAULT 1,
            created_at TEXT DEFAULT '',
            updated_at TEXT DEFAULT '',
            FOREIGN KEY(generation_id) REFERENCES kb_generations(id) ON DELETE CASCADE
        )
        """)
        for column, definition in {
            "generation_id": "INTEGER",
            "source_hash": "TEXT DEFAULT ''",
            "chunking_version": "TEXT DEFAULT ''",
            "parser_name": "TEXT DEFAULT ''",
            "parser_version": "TEXT DEFAULT ''",
            "parse_status": "TEXT DEFAULT 'not_parsed'",
            "page_count": "INTEGER DEFAULT 0",
            "native_page_count": "INTEGER DEFAULT 0",
            "ocr_page_count": "INTEGER DEFAULT 0",
            "needs_ocr_page_count": "INTEGER DEFAULT 0",
            "parse_warning_count": "INTEGER DEFAULT 0",
            "table_count": "INTEGER DEFAULT 0",
            "extracted_image_count": "INTEGER DEFAULT 0",
            "parse_report_json": "TEXT DEFAULT '{}'",
        }.items():
            ensure_column(conn, "documents", column, definition)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_generation ON documents(generation_id)")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS parent_chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_id INTEGER NOT NULL,
            parent_index INTEGER NOT NULL,
            title_path TEXT DEFAULT '',
            content TEXT NOT NULL,
            created_at TEXT DEFAULT '',
            UNIQUE(doc_id, parent_index),
            FOREIGN KEY(doc_id) REFERENCES documents(id) ON DELETE CASCADE
        )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_parent_chunks_doc ON parent_chunks(doc_id, parent_index)")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS chunks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_id INTEGER NOT NULL,
            parent_id INTEGER,
            chunk_index INTEGER NOT NULL,
            content TEXT NOT NULL,
            created_at TEXT DEFAULT '',
            FOREIGN KEY(doc_id) REFERENCES documents(id) ON DELETE CASCADE,
            FOREIGN KEY(parent_id) REFERENCES parent_chunks(id) ON DELETE CASCADE
        )
        """)
        ensure_column(conn, "chunks", "parent_id", "INTEGER")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_chunks_parent ON chunks(parent_id)")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS chunk_embeddings (
            chunk_id INTEGER PRIMARY KEY,
            content_hash TEXT NOT NULL,
            model TEXT NOT NULL,
            dimensions INTEGER NOT NULL,
            vector BLOB NOT NULL,
            updated_at TEXT DEFAULT '',
            FOREIGN KEY(chunk_id) REFERENCES chunks(id) ON DELETE CASCADE
        )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_chunk_embeddings_model ON chunk_embeddings(model)")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS document_pages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_id INTEGER NOT NULL,
            page_number INTEGER NOT NULL,
            parse_method TEXT NOT NULL DEFAULT '',
            parse_status TEXT NOT NULL DEFAULT '',
            text_char_count INTEGER DEFAULT 0,
            meaningful_char_count INTEGER DEFAULT 0,
            quality_score REAL DEFAULT 0,
            garbled_ratio REAL DEFAULT 0,
            image_count INTEGER DEFAULT 0,
            needs_ocr INTEGER DEFAULT 0,
            ocr_confidence REAL,
            warning TEXT DEFAULT '',
            created_at TEXT DEFAULT '',
            UNIQUE(doc_id, page_number),
            FOREIGN KEY(doc_id) REFERENCES documents(id) ON DELETE CASCADE
        )
        """)
        ensure_column(conn, "document_pages", "table_count", "INTEGER DEFAULT 0")
        ensure_column(conn, "document_pages", "extracted_image_count", "INTEGER DEFAULT 0")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_document_pages_doc ON document_pages(doc_id, page_number)")
        conn.execute("""
        CREATE TABLE IF NOT EXISTS document_tables (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_id INTEGER NOT NULL,
            page_number INTEGER NOT NULL,
            table_index INTEGER NOT NULL,
            row_count INTEGER NOT NULL DEFAULT 0,
            column_count INTEGER NOT NULL DEFAULT 0,
            markdown TEXT NOT NULL DEFAULT '',
            rows_json TEXT NOT NULL DEFAULT '[]',
            created_at TEXT DEFAULT '',
            UNIQUE(doc_id, page_number, table_index),
            FOREIGN KEY(doc_id) REFERENCES documents(id) ON DELETE CASCADE
        )
        """)
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_document_tables_doc "
            "ON document_tables(doc_id, page_number, table_index)"
        )
        try:
            ensure_bm25_fts_table(conn)
        except sqlite3.OperationalError:
            # Some local SQLite builds do not include FTS5. The LIKE path below
            # keeps the knowledge base usable in that environment.
            pass
        active_row = conn.execute(
            "SELECT id FROM kb_generations WHERE status = 'active' LIMIT 1"
        ).fetchone()
        if active_row is None:
            cursor = conn.execute(
                """
                INSERT INTO kb_generations(
                    name, status, created_at, updated_at, published_at, metadata_json
                ) VALUES (?, 'active', ?, ?, ?, ?)
                """,
                (
                    "legacy-current",
                    now_str(),
                    now_str(),
                    now_str(),
                    json.dumps({"migration": "existing_data"}, ensure_ascii=False),
                ),
            )
            active_id = int(cursor.lastrowid)
        else:
            active_id = int(active_row["id"])
        conn.execute(
            "UPDATE documents SET generation_id = ? WHERE generation_id IS NULL",
            (active_id,),
        )
        conn.commit()


def get_active_generation_id() -> int:
    init_db()
    with connect() as conn:
        row = conn.execute(
            "SELECT id FROM kb_generations WHERE status = 'active' LIMIT 1"
        ).fetchone()
    if row is None:
        raise RuntimeError("知识库没有 active generation")
    return int(row["id"])


def get_active_generation() -> Dict[str, Any]:
    init_db()
    with connect() as conn:
        row = conn.execute(
            "SELECT * FROM kb_generations WHERE status = 'active' LIMIT 1"
        ).fetchone()
    return dict(row) if row else {}


def normalize_text(text: str) -> str:
    text = text or ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_file_text(path: Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".py", ".yaml", ".yml", ".json", ".html", ".htm"}:
        return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".csv":
        try:
            lines = []
            with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
                reader = csv.reader(f)
                for row in reader:
                    lines.append(" | ".join(row))
            return "\n".join(lines)
        except Exception:
            return path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".docx":
        try:
            import docx
            doc = docx.Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as e:
            return f"DOCX 读取失败：{e}"

    if suffix == ".pdf":
        from rag.pdf_parser import parse_pdf

        return parse_pdf(path).text

    return path.read_text(encoding="utf-8", errors="ignore")


def split_chunks(text: str, chunk_size: int = 700, overlap: int = 120) -> List[str]:
    text = normalize_text(text)
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    # 优先按 Markdown 标题、空行和 PDF 页码边界分段，减少章节语义被硬切断。
    blocks = re.split(r"(?=^#{1,4}\s)|(?=^\[PDF_PAGE=\d+\]$)|\n{2,}", text, flags=re.MULTILINE)
    blocks = [block.strip() for block in blocks if block and block.strip()]
    chunks: List[str] = []
    current = ""

    def flush() -> None:
        nonlocal current
        value = current.strip()
        if value:
            chunks.append(value)
        current = ""

    for block in blocks:
        first_line = block.splitlines()[0] if block else ""
        is_structured_table = first_line.startswith("### 第") and "页表格" in first_line
        if is_structured_table and len(block) <= 5000:
            flush()
            chunks.append(block)
            continue
        if len(block) > chunk_size:
            flush()
            start = 0
            while start < len(block):
                end = min(start + chunk_size, len(block))
                value = block[start:end].strip()
                if value:
                    chunks.append(value)
                if end >= len(block):
                    break
                start = max(0, end - overlap)
            continue

        candidate = f"{current}\n\n{block}".strip() if current else block
        if len(candidate) <= chunk_size:
            current = candidate
        else:
            previous_tail = current[-overlap:].strip() if current and overlap > 0 else ""
            flush()
            current = f"{previous_tail}\n\n{block}".strip() if previous_tail else block

    flush()
    return chunks


def split_parent_child_chunks(
    text: str,
    parent_size: Optional[int] = None,
    child_size: Optional[int] = None,
    child_overlap: Optional[int] = None,
) -> List[Dict[str, Any]]:
    """先按章节形成父块，再在每个父块内部生成可检索子块。"""
    parent_size = int(parent_size or getattr(settings, "PARENT_CHUNK_SIZE", 2000))
    child_size = int(child_size or getattr(settings, "CHILD_CHUNK_SIZE", 500))
    child_overlap = int(child_overlap if child_overlap is not None else getattr(settings, "CHILD_CHUNK_OVERLAP", 80))
    text = normalize_text(text)
    if not text:
        return []

    sections: List[Dict[str, str]] = []
    current_title = "文档正文"
    level_one_title = ""
    current_lines: List[str] = []

    def flush_section() -> None:
        nonlocal current_lines
        value = "\n".join(current_lines).strip()
        meaningful = [
            line for line in value.splitlines()
            if line.strip() and not re.match(r"^#{1,2}\s+", line.strip())
        ]
        if value and meaningful:
            sections.append({"title_path": current_title, "content": value})
        current_lines = []

    for line in text.splitlines():
        stripped = line.strip()
        heading = re.match(r"^(#{1,2})\s+(.+)$", stripped)
        page = re.match(r"^\[PDF_PAGE=(\d+)\]$", stripped)
        if heading or page:
            flush_section()
            if heading:
                heading_text = heading.group(2).strip()
                if len(heading.group(1)) == 1:
                    level_one_title = heading_text
                    current_title = heading_text
                else:
                    current_title = (
                        f"{level_one_title} / {heading_text}"
                        if level_one_title else heading_text
                    )
            else:
                current_title = f"第{page.group(1)}页"
        current_lines.append(line)
    flush_section()

    hierarchy: List[Dict[str, Any]] = []
    parent_index = 0
    for section in sections:
        parent_parts = split_chunks(section["content"], chunk_size=parent_size, overlap=0)
        for part_index, parent_content in enumerate(parent_parts, start=1):
            title_path = section["title_path"]
            if len(parent_parts) > 1:
                title_path = f"{title_path} / 第{part_index}部分"
            child_contents = split_chunks(
                parent_content,
                chunk_size=child_size,
                overlap=child_overlap,
            )
            children = []
            for child_content in child_contents:
                section_prefix = f"[SECTION={title_path}]"
                indexed_content = (
                    child_content if child_content.startswith(section_prefix)
                    else f"{section_prefix}\n{child_content}"
                )
                children.append(indexed_content)
            hierarchy.append({
                "parent_index": parent_index,
                "title_path": title_path,
                "content": parent_content,
                "children": children,
            })
            parent_index += 1
    return hierarchy


def _insert_chunk_hierarchy(
    conn: sqlite3.Connection,
    *,
    doc_id: int,
    title: str,
    category: str,
    content: str,
) -> int:
    hierarchy = split_parent_child_chunks(content)
    child_index = 0
    for parent in hierarchy:
        parent_cursor = conn.execute(
            """
            INSERT INTO parent_chunks(
                doc_id, parent_index, title_path, content, created_at
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (
                doc_id,
                parent["parent_index"],
                parent["title_path"],
                parent["content"],
                now_str(),
            ),
        )
        parent_id = int(parent_cursor.lastrowid)
        for child in parent["children"]:
            child_cursor = conn.execute(
                """
                INSERT INTO chunks(
                    doc_id, parent_id, chunk_index, content, created_at
                ) VALUES (?, ?, ?, ?, ?)
                """,
                (doc_id, parent_id, child_index, child, now_str()),
            )
            chunk_id = int(child_cursor.lastrowid)
            if has_fts_table(conn):
                conn.execute(
                    "INSERT INTO chunks_fts(title, category, content, doc_id, chunk_id) VALUES (?, ?, ?, ?, ?)",
                    (title, category, child, doc_id, chunk_id),
                )
            child_index += 1
    return child_index


def add_text_document(
    title: str,
    content: str,
    category: str = "手动添加",
    source: str = "manual",
    parse_report: Optional[Dict[str, Any]] = None,
    source_hash: str = "",
    generation_id: Optional[int] = None,
) -> int:
    init_db()
    content = normalize_text(content)
    hierarchy = split_parent_child_chunks(content)
    if not hierarchy:
        raise ValueError("内容为空，无法入库")

    with connect() as conn:
        if generation_id is None:
            active_row = conn.execute(
                "SELECT id FROM kb_generations WHERE status = 'active' LIMIT 1"
            ).fetchone()
            if active_row is None:
                raise RuntimeError("知识库没有 active generation")
            generation_id = int(active_row["id"])
        cur = conn.execute(
            """
            INSERT INTO documents(
                generation_id, title, source, category, enabled, created_at, updated_at
            ) VALUES (?, ?, ?, ?, 1, ?, ?)
            """,
            (generation_id, title, source, category, now_str(), now_str()),
        )
        doc_id = int(cur.lastrowid)
        if source_hash:
            conn.execute(
                "UPDATE documents SET source_hash = ? WHERE id = ?",
                (source_hash, doc_id),
            )
        if parse_report:
            conn.execute(
                """
                UPDATE documents SET
                    parser_name = ?, parser_version = ?, parse_status = ?,
                    page_count = ?, native_page_count = ?, ocr_page_count = ?,
                    needs_ocr_page_count = ?, parse_warning_count = ?,
                    table_count = ?, extracted_image_count = ?,
                    parse_report_json = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    parse_report.get("parser_name", ""),
                    parse_report.get("parser_version", ""),
                    parse_report.get("status", "completed"),
                    int(parse_report.get("page_count", 0)),
                    int(parse_report.get("native_page_count", 0)),
                    int(parse_report.get("ocr_page_count", 0)),
                    int(parse_report.get("needs_ocr_page_count", 0)),
                    int(parse_report.get("warning_count", 0)),
                    int(parse_report.get("metadata", {}).get("table_count", 0)),
                    int(parse_report.get("metadata", {}).get("extracted_image_count", 0)),
                    json.dumps(parse_report, ensure_ascii=False),
                    now_str(),
                    doc_id,
                ),
            )
            for page in parse_report.get("pages", []):
                conn.execute(
                    """
                    INSERT INTO document_pages(
                        doc_id, page_number, parse_method, parse_status,
                        text_char_count, meaningful_char_count, quality_score,
                        garbled_ratio, image_count, needs_ocr, ocr_confidence,
                        warning, table_count, extracted_image_count, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        doc_id,
                        int(page.get("page_number", 0)),
                        page.get("parse_method", ""),
                        page.get("status", ""),
                        int(page.get("text_char_count", 0)),
                        int(page.get("meaningful_char_count", 0)),
                        float(page.get("quality_score", 0)),
                        float(page.get("garbled_ratio", 0)),
                        int(page.get("image_count", 0)),
                        int(bool(page.get("needs_ocr", False))),
                        page.get("ocr_confidence"),
                        page.get("warning", ""),
                        int(page.get("table_count", 0)),
                        int(page.get("extracted_image_count", 0)),
                        now_str(),
                    ),
                )
                for table in page.get("tables", []):
                    rows = table.get("rows", [])
                    column_count = max((len(row) for row in rows), default=0)
                    conn.execute(
                        """
                        INSERT INTO document_tables(
                            doc_id, page_number, table_index, row_count,
                            column_count, markdown, rows_json, created_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            doc_id,
                            int(page.get("page_number", 0)),
                            int(table.get("index", 0)),
                            len(rows),
                            column_count,
                            table.get("markdown", ""),
                            json.dumps(rows, ensure_ascii=False),
                            now_str(),
                        ),
                    )
        else:
            conn.execute(
                """
                UPDATE documents SET parser_name = 'plain_text', parser_version = '1.0',
                    parse_status = 'completed', updated_at = ? WHERE id = ?
                """,
                (now_str(), doc_id),
            )
        _insert_chunk_hierarchy(
            conn,
            doc_id=doc_id,
            title=title,
            category=category,
            content=content,
        )
        conn.execute(
            "UPDATE documents SET chunking_version = ? WHERE id = ?",
            (getattr(settings, "CHUNKING_VERSION", "parent_child_v1"), doc_id),
        )
        conn.commit()
    return doc_id


def add_file_document(
    path: Path,
    title: Optional[str] = None,
    category: str = "文件导入",
    generation_id: Optional[int] = None,
) -> int:
    path = Path(path)
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    if path.suffix.lower() == ".pdf":
        from rag.pdf_parser import parse_pdf

        parsed = parse_pdf(path)
        return add_text_document(
            title or path.name,
            parsed.text,
            category=category,
            source=str(path),
            parse_report=parsed.public_dict(include_page_text=False),
            source_hash=source_hash,
            generation_id=generation_id,
        )
    text = read_file_text(path)
    return add_text_document(
        title or path.name,
        text,
        category=category,
        source=str(path),
        source_hash=source_hash,
        generation_id=generation_id,
    )


def import_folder_incremental_report(folder: Path) -> Dict[str, Any]:
    """增量导入目录：相同哈希跳过，内容变化则安全替换旧文档。"""
    init_db()
    folder = Path(folder).resolve()
    imported: List[Dict[str, Any]] = []
    updated: List[Dict[str, Any]] = []
    skipped: List[Dict[str, Any]] = []
    errors: List[Dict[str, str]] = []
    if not folder.exists():
        return {
            "folder": str(folder),
            "imported": imported,
            "updated": updated,
            "skipped": skipped,
            "errors": [{"file": str(folder), "error": "目录不存在"}],
        }

    files = [
        path for path in sorted(folder.rglob("*"))
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTS
    ]
    for path in files:
        resolved_source = str(path.resolve())
        file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        with connect() as conn:
            old_row = conn.execute(
                """
                SELECT id, source_hash, chunking_version
                FROM documents
                WHERE LOWER(source) = LOWER(?)
                ORDER BY id DESC LIMIT 1
                """,
                (resolved_source,),
            ).fetchone()
        current_chunking = getattr(settings, "CHUNKING_VERSION", "parent_child_v1")
        if (
            old_row
            and old_row["source_hash"] == file_hash
            and old_row["chunking_version"] == current_chunking
        ):
            skipped.append({"file": resolved_source, "doc_id": int(old_row["id"])})
            continue
        if old_row and old_row["source_hash"] == file_hash:
            try:
                rechunk_document(int(old_row["id"]))
                updated.append({"file": resolved_source, "doc_id": int(old_row["id"])})
            except Exception as exc:
                errors.append({"file": resolved_source, "error": str(exc)})
            continue
        try:
            new_doc_id = add_file_document(
                path,
                title=path.name,
                category=infer_category(path.name),
            )
            if old_row:
                delete_document(int(old_row["id"]))
                updated.append({"file": resolved_source, "doc_id": new_doc_id})
            else:
                imported.append({"file": resolved_source, "doc_id": new_doc_id})
        except Exception as exc:
            errors.append({"file": resolved_source, "error": str(exc)})

    return {
        "folder": str(folder),
        "files_seen": len(files),
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }


def clear_kb() -> None:
    init_db()
    with connect() as conn:
        if has_fts_table(conn):
            conn.execute("DELETE FROM chunks_fts")
        conn.execute("DELETE FROM chunks")
        conn.execute("DELETE FROM documents")
        conn.commit()


def rebuild_from_folder_report(folder: Path) -> Dict[str, Any]:
    init_db()
    folder = Path(folder)
    clear_kb()
    imported: List[Dict[str, Any]] = []
    errors: List[Dict[str, str]] = []
    if not folder.exists():
        return {
            "success": False,
            "files_seen": 0,
            "imported": [],
            "errors": [{"file": str(folder), "error": "知识库目录不存在"}],
            "chunk_count": 0,
        }
    files = [
        path for path in sorted(folder.rglob("*"))
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTS
    ]
    for path in files:
        try:
            doc_id = add_file_document(path, title=path.name, category=infer_category(path.name))
            report = get_document_parse_report(doc_id)
            imported.append({
                "file": str(path),
                "doc_id": doc_id,
                "parse_status": (report or {}).get("parse_status", "completed"),
                "page_count": (report or {}).get("page_count", 0),
                "needs_ocr_page_count": (report or {}).get("needs_ocr_page_count", 0),
            })
        except Exception as exc:
            errors.append({"file": str(path), "error": str(exc)})
    return {
        "success": not errors,
        "files_seen": len(files),
        "imported": imported,
        "errors": errors,
        "chunk_count": get_status()["chunk_count"],
    }


def rebuild_from_folder(folder: Path) -> int:
    """Backward-compatible wrapper for older callers."""
    return int(rebuild_from_folder_report(folder)["chunk_count"])


def infer_category(name: str) -> str:
    n = name.lower()
    if "双模态" in name or ("红外" in name and "可见光" in name):
        return "红外可见光双模态设备"
    if "纯红外" in name or "红外机芯" in name or "测温机芯" in name:
        return "纯红外设备"
    if "纯可见光" in name or "可见光工业相机" in name:
        return "纯可见光设备"
    if "验收" in name or "sop" in n or "采购" in name or "交付" in name:
        return "企业流程"
    if "标准" in name or "元数据" in name or "stac" in n or "geotiff" in n:
        return "数据治理"
    if "外部" in name or "索引" in name:
        return "外部资料"
    if "产品图" in name or "图谱" in name:
        return "产品图谱"
    if "指标" in name or "qabf" in n or "ssim" in n:
        return "融合指标"
    if "故障" in name or "fault" in n:
        return "故障案例"
    if "维护" in name or "天气" in name or "保养" in name:
        return "维护手册"
    if "选购" in name or "指南" in name:
        return "选购指南"
    if "sar" in n or "雷达" in name:
        return "SAR设备资料"
    if "landsat" in n or "遥感" in name:
        return "遥感数据资料"
    if "光学" in name or "optical" in n:
        return "光学设备资料"
    return "通用知识"


def list_documents() -> List[Dict[str, Any]]:
    init_db()
    with connect() as conn:
        rows = conn.execute("""
            SELECT d.*,
                   COUNT(c.id) AS chunk_count
            FROM documents d
            JOIN kb_generations g ON g.id = d.generation_id AND g.status = 'active'
            LEFT JOIN chunks c ON d.id = c.doc_id
            GROUP BY d.id
            ORDER BY d.id DESC
        """).fetchall()
    documents = [dict(r) for r in rows]
    for document in documents:
        document.pop("parse_report_json", None)
    return documents


def get_document_parse_report(doc_id: int) -> Optional[Dict[str, Any]]:
    init_db()
    with connect() as conn:
        document = conn.execute(
            """
            SELECT id, title, source, parser_name, parser_version, parse_status,
                   page_count, native_page_count, ocr_page_count,
                   needs_ocr_page_count, parse_warning_count, table_count,
                   extracted_image_count, parse_report_json
            FROM documents WHERE id = ?
            """,
            (doc_id,),
        ).fetchone()
        if document is None:
            return None
        pages = conn.execute(
            """
            SELECT page_number, parse_method, parse_status AS status,
                   text_char_count, meaningful_char_count, quality_score,
                   garbled_ratio, image_count, needs_ocr, ocr_confidence, warning
                   , table_count, extracted_image_count
            FROM document_pages WHERE doc_id = ? ORDER BY page_number
            """,
            (doc_id,),
        ).fetchall()
    payload = dict(document)
    raw_report = payload.pop("parse_report_json", "")
    try:
        stored = json.loads(raw_report or "{}")
    except json.JSONDecodeError:
        stored = {}
    payload["warnings"] = stored.get("warnings", [])
    payload["metadata"] = stored.get("metadata", {})
    stored_pages = {
        int(page.get("page_number", 0)): page for page in stored.get("pages", [])
    }
    payload["pages"] = []
    for page in pages:
        row = dict(page)
        merged = {**stored_pages.get(int(row["page_number"]), {}), **row}
        merged["needs_ocr"] = bool(row["needs_ocr"])
        payload["pages"].append(merged)
    return payload


def list_chunks(doc_id: Optional[int] = None, limit: int = 200) -> List[Dict[str, Any]]:
    init_db()
    with connect() as conn:
        if doc_id is None:
            rows = conn.execute("""
                SELECT c.id, c.doc_id, c.parent_id, c.chunk_index, c.content,
                       d.title, d.category, d.enabled
                FROM chunks c JOIN documents d ON c.doc_id = d.id
                JOIN kb_generations g ON g.id = d.generation_id AND g.status = 'active'
                ORDER BY c.id DESC LIMIT ?
            """, (limit,)).fetchall()
        else:
            rows = conn.execute("""
                SELECT c.id, c.doc_id, c.parent_id, c.chunk_index, c.content,
                       d.title, d.category, d.enabled
                FROM chunks c JOIN documents d ON c.doc_id = d.id
                JOIN kb_generations g ON g.id = d.generation_id AND g.status = 'active'
                WHERE c.doc_id = ?
                ORDER BY c.chunk_index ASC LIMIT ?
            """, (doc_id, limit)).fetchall()
    return [dict(r) for r in rows]


def get_parent_chunks(
    parent_ids: List[int], generation_id: Optional[int] = None
) -> List[Dict[str, Any]]:
    init_db()
    ids = sorted({int(value) for value in parent_ids if value is not None})
    if not ids:
        return []
    placeholders = ",".join("?" for _ in ids)
    with connect() as conn:
        generation_sql = (
            "d.generation_id = ?" if generation_id is not None
            else "g.status = 'active'"
        )
        params: List[Any] = list(ids)
        if generation_id is not None:
            params.append(int(generation_id))
        rows = conn.execute(
            f"""
            SELECT p.id AS parent_id, p.doc_id, p.parent_index, p.title_path,
                   p.content, d.title AS doc_title, d.category, d.source
            FROM parent_chunks p
            JOIN documents d ON d.id = p.doc_id
            JOIN kb_generations g ON g.id = d.generation_id
            WHERE d.enabled = 1 AND p.id IN ({placeholders}) AND {generation_sql}
            """,
            params,
        ).fetchall()
    return [dict(row) for row in rows]


def rechunk_document(doc_id: int) -> Dict[str, Any]:
    """保留文档ID和解析报告，仅重建父块、子块、FTS及其后续向量。"""
    init_db()
    with connect() as conn:
        document = conn.execute(
            "SELECT id, title, source, category FROM documents WHERE id = ?",
            (doc_id,),
        ).fetchone()
        if document is None:
            raise ValueError(f"文档不存在：{doc_id}")
        old_chunks = conn.execute(
            "SELECT content FROM chunks WHERE doc_id = ? ORDER BY chunk_index",
            (doc_id,),
        ).fetchall()

    source_path = Path(str(document["source"] or ""))
    if source_path.exists() and source_path.is_file():
        content = read_file_text(source_path)
    else:
        content = "\n\n".join(row["content"] for row in old_chunks)
    if not normalize_text(content):
        raise ValueError(f"文档内容为空：{document['title']}")

    with connect() as conn:
        if has_fts_table(conn):
            conn.execute("DELETE FROM chunks_fts WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM parent_chunks WHERE doc_id = ?", (doc_id,))
        child_count = _insert_chunk_hierarchy(
            conn,
            doc_id=doc_id,
            title=document["title"],
            category=document["category"],
            content=content,
        )
        parent_count = int(conn.execute(
            "SELECT COUNT(*) FROM parent_chunks WHERE doc_id = ?", (doc_id,)
        ).fetchone()[0])
        conn.execute(
            "UPDATE documents SET chunking_version = ?, updated_at = ? WHERE id = ?",
            (getattr(settings, "CHUNKING_VERSION", "parent_child_v1"), now_str(), doc_id),
        )
        conn.commit()
    return {
        "doc_id": doc_id,
        "title": document["title"],
        "parent_count": parent_count,
        "child_count": child_count,
    }


def rechunk_all_documents() -> Dict[str, Any]:
    init_db()
    with connect() as conn:
        doc_ids = [int(row[0]) for row in conn.execute(
            """
            SELECT d.id FROM documents d
            JOIN kb_generations g ON g.id = d.generation_id
            WHERE g.status = 'active'
            ORDER BY d.id
            """
        ).fetchall()]
    completed = []
    errors = []
    for doc_id in doc_ids:
        try:
            completed.append(rechunk_document(doc_id))
        except Exception as exc:
            errors.append({"doc_id": doc_id, "error": str(exc)})
    return {"completed": completed, "errors": errors}


def delete_document(doc_id: int) -> None:
    init_db()
    with connect() as conn:
        if has_fts_table(conn):
            conn.execute("DELETE FROM chunks_fts WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        conn.commit()


def set_document_enabled(doc_id: int, enabled: bool) -> None:
    init_db()
    with connect() as conn:
        conn.execute("UPDATE documents SET enabled = ?, updated_at = ? WHERE id = ?", (1 if enabled else 0, now_str(), doc_id))
        conn.commit()


SYNONYM_GROUPS = [
    {"模糊", "不清楚", "不清晰", "清晰度低", "发虚", "虚焦", "失焦", "对焦不准", "边缘发虚"},
    {"噪声", "噪点", "颗粒", "斑点", "相干斑", "杂波"},
    {"错位", "偏移", "配准", "对不齐", "重影", "结构漂移"},
    {"进水", "受潮", "潮湿", "水汽", "凝露"},
    {"低照度", "夜间", "光线暗", "亮度低", "照度不足"},
    {"产品图", "产品图片", "设备图片", "外观图", "产品图谱"},
    {"验收", "交付", "放行", "签收", "质量门"},
    {"元数据", "数据目录", "资产目录", "STAC", "数据血缘"},
]


def ordered_unique(items: List[str], limit: int = 24) -> List[str]:
    out: List[str] = []
    seen = set()
    for item in items:
        value = (item or "").strip()
        if len(value) < 2:
            continue
        key = value.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(value)
        if len(out) >= limit:
            break
    return out


def chinese_ngrams(text: str) -> List[str]:
    compact = re.sub(r"[\s,，。；;：:、/\\|]+", "", text or "")
    grams: List[str] = []
    for size in [4, 3, 2]:
        for i in range(0, max(len(compact) - size + 1, 0)):
            gram = compact[i:i + size]
            if any(ch in gram for ch in ["吗", "呢", "的", "了", "是", "么", "什"]):
                continue
            grams.append(gram)
    return grams[:16]


def tokenize_query(query: str) -> List[str]:
    raw = re.split(r"[\s,，。；;：:、/\\|]+", query.strip())
    terms = [t.strip() for t in raw if len(t.strip()) >= 2]
    # 对无空格中文查询，额外保留完整句和关键子串
    if query.strip() and query.strip() not in terms:
        terms.insert(0, query.strip())
    keys = [
        "Qabf", "SSIM", "MI", "SAR", "光学", "融合", "图像", "清晰", "模糊", "不清楚",
        "噪声", "维护", "海边", "雨天", "边缘", "对焦", "镜头", "配准",
        "SLC", "GRD", "GeoTIFF", "COG", "STAC", "产品", "图片", "交付", "验收",
    ]
    for k in keys:
        if k.lower() in query.lower() and k not in terms:
            terms.append(k)
    for group in SYNONYM_GROUPS:
        if any(word.lower() in query.lower() for word in group):
            terms.extend(sorted(group, key=len, reverse=True))
    terms.extend(chinese_ngrams(query))
    return ordered_unique(terms)


def build_fts_match_expression(terms: List[str]) -> str:
    """Build a safely quoted OR query for FTS5 MATCH."""
    expressions: List[str] = []
    for term in terms:
        value = (term or "").strip()
        # The trigram tokenizer cannot match one- or two-character terms.
        if len(value) < 3:
            continue
        escaped = value.replace('"', '""')
        expressions.append(f'"{escaped}"')
    return " OR ".join(expressions[:24])


def search_knowledge(
    query: str, top_k: int = 5, generation_id: Optional[int] = None
) -> Dict[str, Any]:
    init_db()
    query = (query or "").strip()
    if not query:
        return {"confidence": "low", "confidence_reason": "query 为空", "top_score": 0, "sources": [], "context": "", "source_text": ""}

    terms = tokenize_query(query)
    sources: List[Dict[str, Any]] = []

    retrieval_method = "like_fallback"
    with connect() as conn:
        generation_sql = (
            "d.generation_id = ?" if generation_id is not None
            else "g.status = 'active'"
        )
        generation_params = [int(generation_id)] if generation_id is not None else []
        rows = []

        # 1) FTS5倒排召回 + SQLite原生BM25排序。标题和类别权重
        # 高于正文；SQLite返回值越小（越负），相关性越高。
        match_expr = build_fts_match_expression(terms)
        if match_expr and has_fts_table(conn):
            rows = conn.execute(f"""
                SELECT c.id AS chunk_id, c.doc_id, c.parent_id, c.chunk_index,
                       c.content, d.title, d.category, d.source, d.enabled,
                       bm25(chunks_fts, 8.0, 4.0, 1.0) AS bm25_rank
                FROM chunks_fts
                JOIN documents d ON chunks_fts.doc_id = d.id
                JOIN chunks c ON chunks_fts.chunk_id = c.id
                JOIN kb_generations g ON g.id = d.generation_id
                WHERE d.enabled = 1 AND chunks_fts MATCH ? AND {generation_sql}
                ORDER BY bm25_rank ASC
                LIMIT 80
            """, [match_expr, *generation_params]).fetchall()
            if rows:
                retrieval_method = "fts5_bm25"

        # 2) FTS5不可用或未命中时，LIKE保证两字中文词仍能召回。
        if not rows:
            where_parts = []
            params: List[Any] = []
            for term in terms:
                where_parts.append("(c.content LIKE ? OR d.title LIKE ? OR d.category LIKE ?)")
                like = f"%{term}%"
                params.extend([like, like, like])

            where_sql = " OR ".join(where_parts) if where_parts else "c.content LIKE ?"
            if not params:
                params.append(f"%{query}%")
            rows = conn.execute(f"""
                SELECT c.id AS chunk_id, c.doc_id, c.parent_id, c.chunk_index,
                       c.content, d.title, d.category, d.source, d.enabled,
                       NULL AS bm25_rank
                FROM chunks c JOIN documents d ON c.doc_id = d.id
                JOIN kb_generations g ON g.id = d.generation_id
                WHERE d.enabled = 1 AND ({where_sql}) AND {generation_sql}
                LIMIT 2000
            """, params + generation_params).fetchall()

    bm25_values = [
        max(-float(row["bm25_rank"]), 0.0)
        for row in rows if row["bm25_rank"] is not None
    ]
    best_bm25 = max(bm25_values, default=0.0)

    for row_index, r in enumerate(rows):
        d = dict(r)
        content = d.get("content", "") or ""
        raw_bm25 = d.get("bm25_rank")
        if raw_bm25 is not None:
            positive_bm25 = max(-float(raw_bm25), 0.0)
            score = positive_bm25 / best_bm25 if best_bm25 > 0 else 1.0 / (row_index + 1)
        else:
            score = score_text(query, terms, content, d.get("title", ""), d.get("category", ""))
        sources.append({
            "doc_id": d.get("doc_id"),
            "chunk_id": d.get("chunk_id"),
            "parent_id": d.get("parent_id"),
            "doc_title": d.get("title", "未知文档"),
            "category": d.get("category", ""),
            "source": d.get("source", ""),
            "chunk_index": d.get("chunk_index", 0),
            "final_score": round(score, 4),
            "sparse_retriever": retrieval_method,
            "bm25_raw_score": round(float(raw_bm25), 8) if raw_bm25 is not None else None,
            "text": content[:800],
        })

    sources.sort(key=lambda x: x["final_score"], reverse=True)

    # 避免一份大型PDF占满全部结果，优先保留文档多样性。
    diversified: List[Dict[str, Any]] = []
    per_doc: Dict[Any, int] = {}
    for source in sources:
        doc_id = source.get("doc_id")
        # 表格与正文现在会分别成块；每份文档最多保留 4 个证据块，
        # 避免只返回产品标题附近的块而漏掉参数表。
        if per_doc.get(doc_id, 0) >= 4:
            continue
        diversified.append(source)
        per_doc[doc_id] = per_doc.get(doc_id, 0) + 1
        if len(diversified) >= top_k:
            break
    sources = diversified
    top_score = sources[0]["final_score"] if sources else 0

    if not sources:
        confidence = "low"
        reason = "未命中知识库内容"
    elif top_score >= 0.45:
        confidence = "high"
        reason = "命中文档与问题关键词匹配度较高"
    elif top_score >= 0.20:
        confidence = "medium"
        reason = "命中文档较少或匹配度一般"
    else:
        confidence = "low"
        reason = "检索分数偏低，建议补充更相关知识"

    context = "\n\n".join([f"【来源{i+1}】{s['doc_title']} | {s['category']}\n{s['text']}" for i, s in enumerate(sources)])
    source_text = "\n".join([f"[{i+1}] {s['doc_title']} | score={s['final_score']}" for i, s in enumerate(sources)])

    return {
        "confidence": confidence,
        "confidence_reason": reason,
        "top_score": top_score,
        "sources": sources,
        "context": context,
        "source_text": source_text,
        "retrieval": {
            "backend": "sqlite_fts5" if retrieval_method == "fts5_bm25" else "sqlite_like",
            "ranker": "bm25" if retrieval_method == "fts5_bm25" else "rule_score",
            "method": retrieval_method,
        },
    }


def score_text(query: str, terms: List[str], content: str, title: str, category: str) -> float:
    text = f"{title} {category} {content}".lower()
    hit = 0
    for t in terms:
        if t.lower() in text:
            hit += 1
    base = hit / max(min(len(terms), 8), 1)
    if query.lower() in text:
        base += 0.25
    title_lower = (title or "").lower()
    if any(term.lower() in title_lower for term in terms):
        base += 0.18
    if any(k in query for k in ["不清楚", "不清晰", "模糊", "发虚", "虚焦", "清晰度"]):
        if any(k in text for k in ["模糊", "不清楚", "不清晰", "发虚", "虚焦", "清晰度", "对焦"]):
            base += 0.30
    if any(k in query for k in ["Qabf", "指标", "SSIM", "MI"]) and "指标" in category:
        base += 0.20
    if any(k in query for k in ["故障", "噪声", "斑点"]) and "故障" in category:
        base += 0.20
    if any(k in query for k in ["维护", "雨", "天气", "进水"]) and "维护" in category:
        base += 0.20
    if any(k in query for k in ["图片", "产品图", "外观"]) and any(k in text for k in ["产品图谱", "产品示意图", "knowledge_assets"]):
        base += 0.35
    if "镜头" in query and any(k in text for k in ["没有普通光学镜头", "不需要普通光学镜头", "射频透波面"]):
        base += 0.45
    if any(k in query.upper() for k in ["SLC", "GRD"]) and "数据产品" in text:
        base += 0.35
    if any(k in query for k in ["交付", "验收", "放行"]) and any(k in category for k in ["企业流程", "验收"]):
        base += 0.30
    if any(k.lower() in query.lower() for k in ["GeoTIFF", "COG", "STAC", "元数据"]) and "数据治理" in category:
        base += 0.35
    return min(base, 1.0)


def get_status() -> Dict[str, Any]:
    init_db()
    with connect() as conn:
        active_row = conn.execute(
            "SELECT id, name FROM kb_generations WHERE status = 'active' LIMIT 1"
        ).fetchone()
        active_id = int(active_row["id"]) if active_row else -1
        active_name = active_row["name"] if active_row else ""
        doc_count = conn.execute(
            "SELECT COUNT(*) FROM documents WHERE generation_id = ?", (active_id,)
        ).fetchone()[0]
        enabled_doc_count = conn.execute(
            "SELECT COUNT(*) FROM documents WHERE generation_id = ? AND enabled = 1",
            (active_id,),
        ).fetchone()[0]
        chunk_count = conn.execute(
            """
            SELECT COUNT(*) FROM chunks c
            JOIN documents d ON d.id = c.doc_id
            WHERE d.generation_id = ?
            """,
            (active_id,),
        ).fetchone()[0]
        parent_chunk_count = conn.execute(
            """
            SELECT COUNT(*) FROM parent_chunks p
            JOIN documents d ON d.id = p.doc_id
            WHERE d.generation_id = ?
            """,
            (active_id,),
        ).fetchone()[0]
        needs_ocr_page_count = conn.execute(
            "SELECT COALESCE(SUM(needs_ocr_page_count), 0) FROM documents WHERE generation_id = ?",
            (active_id,),
        ).fetchone()[0]
        parse_warning_count = conn.execute(
            "SELECT COALESCE(SUM(parse_warning_count), 0) FROM documents WHERE generation_id = ?",
            (active_id,),
        ).fetchone()[0]
        table_count = conn.execute(
            "SELECT COALESCE(SUM(table_count), 0) FROM documents WHERE generation_id = ?",
            (active_id,),
        ).fetchone()[0]
        extracted_image_count = conn.execute(
            "SELECT COALESCE(SUM(extracted_image_count), 0) FROM documents WHERE generation_id = ?",
            (active_id,),
        ).fetchone()[0]
        fts5_available = has_fts_table(conn)
        embedding_count = conn.execute(
            """
            SELECT COUNT(*) FROM chunk_embeddings e
            JOIN chunks c ON c.id = e.chunk_id
            JOIN documents d ON d.id = c.doc_id
            WHERE d.generation_id = ?
            """,
            (active_id,),
        ).fetchone()[0]
        embedding_models = [
            row[0] for row in conn.execute(
                """
                SELECT DISTINCT e.model FROM chunk_embeddings e
                JOIN chunks c ON c.id = e.chunk_id
                JOIN documents d ON d.id = c.doc_id
                WHERE d.generation_id = ? ORDER BY e.model
                """,
                (active_id,),
            ).fetchall()
        ]
    try:
        from rag.light_pdf_adapter import parser_runtime_status

        pdf_parser_runtime = parser_runtime_status()
    except Exception as exc:
        pdf_parser_runtime = {
            "mode": "lightweight",
            "parser_name": "pymupdf_rapidocr_pdfplumber",
            "available": False,
            "error": str(exc),
        }
    return {
        "db_path": str(DB_PATH),
        "active_generation_id": active_id,
        "active_generation": active_name,
        "doc_count": doc_count,
        "enabled_doc_count": enabled_doc_count,
        "chunk_count": chunk_count,
        "parent_chunk_count": parent_chunk_count,
        "needs_ocr_page_count": needs_ocr_page_count,
        "parse_warning_count": parse_warning_count,
        "table_count": table_count,
        "extracted_image_count": extracted_image_count,
        "fts5_available": fts5_available,
        "embedding_count": embedding_count,
        "embedding_models": embedding_models,
        "pdf_parser": pdf_parser_runtime,
    }
